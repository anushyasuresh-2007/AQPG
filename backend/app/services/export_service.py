"""Export service for generating professional DOCX and PDF question papers and answer keys."""

import io
import json
from typing import Any, Dict, List, Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.generated_paper import GeneratedPaper


def generate_paper_docx(paper: GeneratedPaper, include_answers: bool = False) -> io.BytesIO:
    """Generate a clean, styled Microsoft Word (.docx) exam paper."""
    doc = docx.Document()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(f"{paper.board_name.upper()}\n")
    title_run.bold = True
    title_run.font.size = Pt(15)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"{paper.class_name.upper()} EXAMINATION · {paper.subject_name.upper()}\n")
    sub_run.bold = True
    sub_run.font.size = Pt(13)

    if paper.unit_names:
        u_run = sub_p.add_run(f"Units Covered: {paper.unit_names}\n")
        u_run.italic = True
        u_run.font.size = Pt(9.5)
        u_run.font.color.rgb = RGBColor(100, 100, 100)

    # Time and Marks Table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(3.5)

    hdr_cells = table.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("TIME ALLOWED: 3 HOURS").bold = True
    p2 = hdr_cells[1].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.add_run(f"MAXIMUM MARKS: {paper.total_marks}").bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # General Instructions Box
    inst_p = doc.add_paragraph()
    inst_p.paragraph_format.space_after = Pt(10)
    inst_head = inst_p.add_run("General Instructions:\n")
    inst_head.bold = True
    inst_head.font.size = Pt(10)

    instructions = [
        "1. All questions are compulsory.",
        "2. The question paper comprises multiple sections with marks indicated against each question.",
        "3. Write clear, structured steps for numerical and analytical questions.",
        "4. Use diagrams wherever appropriate to substantiate your answers.",
    ]
    for inst in instructions:
        p = doc.add_paragraph(inst, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Questions Grouped by Section
    questions_list = json.loads(paper.questions_data)

    section_a = [q for q in questions_list if q["marks"] <= 1]
    section_b = [q for q in questions_list if q["marks"] == 2]
    section_c = [q for q in questions_list if 3 <= q["marks"] <= 4]
    section_d = [q for q in questions_list if q["marks"] >= 5]

    overall_num = 1
    sections_def = [
        ("SECTION A: Objective & Multiple Choice Questions (1 Mark Each)", section_a),
        ("SECTION B: Very Short Answer Questions (2 Marks Each)", section_b),
        ("SECTION C: Short / Numerical Answer Questions (3-4 Marks Each)", section_c),
        ("SECTION D: Long Answer & Case Study / Application Questions (5+ Marks Each)", section_d),
    ]

    for sec_title, q_group in sections_def:
        if not q_group:
            continue

        sec_hdr = doc.add_paragraph()
        sec_hdr.paragraph_format.space_before = Pt(10)
        sec_hdr.paragraph_format.space_after = Pt(4)
        s_run = sec_hdr.add_run(sec_title)
        s_run.bold = True
        s_run.font.size = Pt(11)

        for q in q_group:
            q_p = doc.add_paragraph()
            q_p.paragraph_format.space_after = Pt(4)
            q_num_run = q_p.add_run(f"Q{overall_num}. ")
            q_num_run.bold = True
            q_num_run.font.size = Pt(10.5)

            q_txt_run = q_p.add_run(f"{q['question']} ")
            q_txt_run.font.size = Pt(10.5)

            m_run = q_p.add_run(f"[{q['marks']} Mark{'s' if q['marks'] > 1 else ''}]")
            m_run.bold = True
            m_run.font.size = Pt(10)

            if include_answers and q.get("answer"):
                ans_p = doc.add_paragraph()
                ans_p.paragraph_format.left_indent = Inches(0.4)
                ans_p.paragraph_format.space_after = Pt(6)
                a_run = ans_p.add_run(f"Model Answer / Solution:\n{q['answer']}\n")
                a_run.italic = True
                a_run.font.size = Pt(9.5)
                a_run.font.color.rgb = RGBColor(0, 100, 0)

            overall_num += 1

    # Footer Total
    end_p = doc.add_paragraph()
    end_p.paragraph_format.space_before = Pt(14)
    end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end_run = end_p.add_run("--- END OF QUESTION PAPER ---")
    end_run.bold = True
    end_run.font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def generate_paper_pdf(paper: GeneratedPaper, include_answers: bool = False) -> io.BytesIO:
    """Generate clean, publication-ready PDF question paper."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        alignment=1,
        spaceAfter=4,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        alignment=1,
        spaceAfter=4,
        textColor=colors.HexColor("#1e293b"),
    )
    meta_style = ParagraphStyle(
        "DocMeta",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        alignment=1,
        spaceAfter=8,
        textColor=colors.HexColor("#64748b"),
    )
    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=colors.HexColor("#0f172a"),
    )
    section_style = ParagraphStyle(
        "SectionHdr",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10,
        spaceBefore=10,
        spaceAfter=4,
        textColor=colors.HexColor("#1e3a8a"),
    )
    q_style = ParagraphStyle(
        "QuestionText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        spaceAfter=5,
        textColor=colors.HexColor("#0f172a"),
    )
    ans_style = ParagraphStyle(
        "AnswerText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=12,
        leftIndent=15,
        spaceAfter=6,
        textColor=colors.HexColor("#065f46"),
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph(f"{paper.board_name.upper()} EXAMINATION", title_style))
    story.append(Paragraph(f"{paper.class_name.upper()} · SUBJECT: {paper.subject_name.upper()}", subtitle_style))
    if paper.unit_names:
        story.append(Paragraph(f"Units Covered: {paper.unit_names}", meta_style))

    # Time & Maximum Marks Bar
    t_data = [
        [
            Paragraph("TIME ALLOWED: 3 HOURS", table_cell_style),
            Paragraph(f"MAXIMUM MARKS: {paper.total_marks}", ParagraphStyle("RightM", parent=table_cell_style, alignment=2)),
        ]
    ]
    t = Table(t_data, colWidths=[260, 260])
    t.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, -1), 1.5, colors.black),
        ("LINEBELOW", (0, 0), (-1, -1), 1.5, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 8))

    # General instructions
    story.append(Paragraph("<b>General Instructions:</b> All questions are compulsory. Questions carry marks indicated on the right. Write clear, legible, and step-by-step solutions where appropriate.", ParagraphStyle("Inst", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=11, spaceAfter=8, textColor=colors.HexColor("#334155"))))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cbd5e1"), spaceAfter=8))

    questions_list = json.loads(paper.questions_data)
    section_a = [q for q in questions_list if q["marks"] <= 1]
    section_b = [q for q in questions_list if q["marks"] == 2]
    section_c = [q for q in questions_list if 3 <= q["marks"] <= 4]
    section_d = [q for q in questions_list if q["marks"] >= 5]

    overall_num = 1
    sections_def = [
        ("SECTION A: Objective & Multiple Choice Questions (1 Mark Each)", section_a),
        ("SECTION B: Very Short Answer Questions (2 Marks Each)", section_b),
        ("SECTION C: Short / Numerical Answer Questions (3-4 Marks Each)", section_c),
        ("SECTION D: Long Answer & Case Study / Application Questions (5+ Marks Each)", section_d),
    ]

    for sec_title, q_group in sections_def:
        if not q_group:
            continue

        story.append(Paragraph(sec_title, section_style))

        for q in q_group:
            q_html = f"<b>Q{overall_num}.</b> {q['question']} &nbsp;&nbsp;<b>[{q['marks']} Mark{'s' if q['marks'] > 1 else ''}]</b>"
            story.append(Paragraph(q_html, q_style))

            if include_answers and q.get("answer"):
                story.append(Paragraph(f"<b>Answer Key / Solution:</b><br/>{q['answer']}", ans_style))

            overall_num += 1

    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceAfter=6))
    story.append(Paragraph("<b>END OF QUESTION PAPER</b>", ParagraphStyle("End", parent=title_style, fontSize=9, alignment=1)))

    doc.build(story)
    buf.seek(0)
    return buf
